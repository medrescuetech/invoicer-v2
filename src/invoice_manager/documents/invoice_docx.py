"""Generate a Microsoft Word .docx version of an invoice."""

from __future__ import annotations

from decimal import Decimal
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches

from invoice_manager.domain.money import Money
from invoice_manager.persistence.models import Invoice


def _get(settings: dict[str, Any], key: str, default: str = "") -> str:
    value = settings.get(key, default)
    if not value:
        return default
    return str(value)


def _label(settings: dict[str, Any], key: str, is_quote: bool, default: str = "") -> str:
    """Return a quote-specific label if available, otherwise the invoice label."""
    if is_quote:
        quote_key = f"quote_{key.removeprefix('invoice_')}"
        value = settings.get(quote_key)
        if value:
            return str(value)
    return _get(settings, key, default)


def _fmt(cents: int) -> str:
    return Money(cents=cents).__str__()


def _generate_docx(invoice: Invoice, settings: dict[str, Any], output_path: Path, is_quote: bool) -> Path:
    """Create a .docx document for an invoice or quote."""
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    document = Document()

    business = _get(settings, "business_name", "Invoice")
    para = document.add_paragraph()
    run = para.add_run(business)
    run.bold = True
    run.font.size = para.runs[0].font.size
    para.alignment = WD_ALIGN_PARAGRAPH.LEFT

    address = _get(settings, "business_address")
    if address:
        document.add_paragraph(address)

    gst_rate = Decimal(_get(settings, "gst_rate", "0.0") or "0.0")
    if is_quote:
        doc_title = (
            _get(settings, "quote_title_tax", "TAX QUOTE")
            if gst_rate > 0
            else _get(settings, "quote_title", "QUOTE")
        )
        date_label = _get(settings, "quote_date_label", "Quote date:")
        due_label = _get(settings, "quote_due_date_label", "Valid until:")
        due_value = str(invoice.due_date) if invoice.due_date else _get(settings, "quote_due_date_na_text", "N/A")
    else:
        doc_title = (
            _get(settings, "invoice_title_tax", "TAX INVOICE")
            if gst_rate > 0
            else _get(settings, "invoice_title", "INVOICE")
        )
        date_label = _get(settings, "invoice_date_label", "Date:")
        due_label = _get(settings, "invoice_due_date_label", "Due date:")
        due_value = str(invoice.due_date) if invoice.due_date else _get(settings, "invoice_due_date_na_text", "N/A")
    title = document.add_paragraph()
    run = title.add_run(f"{doc_title} — {invoice.number}")
    run.bold = True
    run.font.size = title.runs[0].font.size

    document.add_paragraph()
    meta = [
        (date_label, str(invoice.issue_date)),
        (due_label, due_value),
        (_label(settings, "invoice_client_label", is_quote, "Client:"), invoice.client_name),
        (_label(settings, "invoice_address_label", is_quote, "Address:"), invoice.client_address or ""),
    ]
    for label, value in meta:
        p = document.add_paragraph()
        p.add_run(label).bold = True
        p.add_run(f" {value}")

    table = document.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    table.autofit = False
    table.allow_autofit = False
    widths = [Inches(3.0), Inches(0.55), Inches(0.75), Inches(0.9), Inches(0.75), Inches(0.9)]
    for idx, width in enumerate(widths):
        table.columns[idx].width = width
    hdr_cells = table.rows[0].cells
    headers = [
        _label(settings, "invoice_description_header", is_quote, "Description"),
        _label(settings, "invoice_qty_header", is_quote, "Qty"),
        _label(settings, "invoice_unit_header", is_quote, "Unit"),
        _label(settings, "invoice_price_header", is_quote, "Price"),
        _label(settings, "invoice_gst_header", is_quote, "GST"),
        _label(settings, "invoice_total_header", is_quote, "Total"),
    ]
    for idx, header in enumerate(headers):
        hdr_cells[idx].text = header
        for paragraph in hdr_cells[idx].paragraphs:
            for run in paragraph.runs:
                run.font.bold = True

    for item in invoice.items:
        row_cells = table.add_row().cells
        row_cells[0].text = item.description
        row_cells[1].text = str(item.quantity)
        row_cells[2].text = item.unit or "ea"
        row_cells[3].text = _fmt(item.unit_price_cents)
        row_cells[4].text = _fmt(item.gst_cents)
        row_cells[5].text = _fmt(item.total_cents)

    for label, amount in [
        (_label(settings, "invoice_subtotal_label", is_quote, "Subtotal"), _fmt(invoice.subtotal_cents)),
        (_label(settings, "invoice_gst_label", is_quote, "GST"), _fmt(invoice.gst_cents)),
        (_label(settings, "invoice_total_label", is_quote, "Total"), _fmt(invoice.total_cents)),
    ]:
        row_cells = table.add_row().cells
        row_cells[0].text = ""
        row_cells[1].text = ""
        row_cells[2].text = ""
        row_cells[3].text = label
        row_cells[4].text = ""
        row_cells[5].text = amount
        for paragraph in row_cells[3].paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
        for paragraph in row_cells[5].paragraphs:
            for run in paragraph.runs:
                run.font.bold = True

    document.add_paragraph()

    if not is_quote:
        bank_name = _get(settings, "bank_name")
        bsb = _get(settings, "bank_bsb")
        account = _get(settings, "bank_account")
        account_name = _get(settings, "bank_account_name")
        if bank_name or account:
            p = document.add_paragraph()
            p.add_run(_label(settings, "invoice_payment_details_label", is_quote, "Payment details")).bold = True
            payment_line = (
                f"{_label(settings, 'invoice_bank_label', is_quote, 'Bank:')} {bank_name}  |  "
                f"{_label(settings, 'invoice_bsb_label', is_quote, 'BSB:')} {bsb}  |  "
                f"{_label(settings, 'invoice_account_label', is_quote, 'Account:')} {account}  |  "
                f"{_label(settings, 'invoice_account_name_label', is_quote, 'Name:')} {account_name}"
            )
            document.add_paragraph(payment_line)

    if invoice.notes:
        p = document.add_paragraph()
        p.add_run(_label(settings, "invoice_notes_label", is_quote, "Notes:")).bold = True
        p.add_run(f" {invoice.notes}")

    if is_quote:
        quote_footer = _get(settings, "quote_footer_note", "")
        if quote_footer:
            document.add_paragraph(quote_footer)
    else:
        gst_footer = _get(settings, "invoice_gst_footer_note", "")
        if gst_footer:
            document.add_paragraph(gst_footer)

    thank_you = _label(settings, "invoice_thank_you", is_quote, "Thank you for your business!")
    if thank_you:
        document.add_paragraph(thank_you)

    document.save(str(output_path))
    return output_path


def generate_invoice_docx(invoice: Invoice, settings: dict[str, Any], output_path: Path) -> Path:
    """Create a .docx document for the invoice."""
    return _generate_docx(invoice, settings, output_path, is_quote=False)


def generate_quote_docx(invoice: Invoice, settings: dict[str, Any], output_path: Path) -> Path:
    """Create a .docx document for the quote."""
    return _generate_docx(invoice, settings, output_path, is_quote=True)
